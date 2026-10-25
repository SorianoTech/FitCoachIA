# ruff: noqa: E702, S603, S404, T201, N806
"""Regenera docs/memoria-proyecto.docx desde docs/memoria-proyecto.md.

Uso (desde la raiz del repo):
    uv run --no-project --with pypandoc_binary --with python-docx --with matplotlib \\
        python tools/memoria/build_memoria.py [--sin-diagramas]

Pasos: 1) diagramas PNG en docs/img, 2) pandoc md -> docx con indice,
3) maquetacion (portada, estilos, tablas, cabecera/pie). Word actualiza
los campos (indice, n. de paginas) al abrir el documento.
"""

import subprocess
import sys
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BLUE, GREY = RGBColor(0x1A, 0x3A, 0x6B), RGBColor(0x55, 0x55, 0x55)


def font(style, name="Calibri", size=None, bold=None, color=None, italic=None):
    f = style.font
    f.name = name
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        rf.attrib.pop(qn(a), None)
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color


def fmt(style, before, after, align=None, line=None, keep=False):
    p = style.paragraph_format
    p.space_before, p.space_after = Pt(before), Pt(after)
    if line:
        p.line_spacing = line
    if align is not None:
        p.alignment = align
    if keep:
        p.keep_with_next = True


def style_docx(d):
    S = d.styles
    for n in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if n in [s.name for s in S]:
            font(S[n], size=10.5)
    fmt(S["Normal"], 0, 4, line=1.12)
    for n in ("Body Text", "First Paragraph"):
        fmt(S[n], 0, 6, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.12)
    fmt(S["Compact"], 0, 2, line=1.08)
    font(S["Heading 1"], size=18, bold=True, color=BLUE)
    fmt(S["Heading 1"], 18, 8, keep=True)
    font(S["Heading 2"], size=13.5, bold=True, color=BLUE)
    fmt(S["Heading 2"], 12, 5, keep=True)
    font(S["Heading 3"], size=11.5, bold=True, color=GREY)
    fmt(S["Heading 3"], 8, 3, keep=True)
    font(S["Title"], size=28, bold=True, color=BLUE)
    fmt(S["Title"], 0, 10, WD_ALIGN_PARAGRAPH.CENTER)
    font(S["Subtitle"], size=13, color=GREY, italic=False)
    fmt(S["Subtitle"], 0, 24, WD_ALIGN_PARAGRAPH.CENTER)
    font(S["Author"], size=12, color=BLUE)
    fmt(S["Author"], 0, 6, WD_ALIGN_PARAGRAPH.CENTER)
    for n in ("Block Text",):
        font(S[n], size=9.5, italic=True, color=GREY)
        fmt(S[n], 6, 6, WD_ALIGN_PARAGRAPH.LEFT)
    for n in ("Image Caption", "Captioned Figure", "Figure"):
        if n in [s.name for s in S]:
            font(S[n], size=9, italic=True, color=GREY)
            fmt(S[n], 2, 10, WD_ALIGN_PARAGRAPH.CENTER)
    for n in ("TOC Heading",):
        if n in [s.name for s in S]:
            font(S[n], size=18, bold=True, color=BLUE)

    # Pagina A4 y margenes
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.4)
    sec.top_margin, sec.bottom_margin = Cm(2.3), Cm(2.2)
    sec.different_first_page_header_footer = True

    def field(run, instr):
        for t, txt in (("begin", None), (None, instr), ("end", None)):
            if t:
                e = OxmlElement("w:fldChar")
                e.set(qn("w:fldCharType"), t)
            else:
                e = OxmlElement("w:instrText")
                e.set(qn("xml:space"), "preserve")
                e.text = txt
            run._r.append(e)

    def border(p, side, color="1A3A6B", sz=6):
        ppr = p._p.get_or_add_pPr()
        b = OxmlElement("w:pBdr")
        e = OxmlElement(f"w:{side}")
        for k, v in (("val", "single"), ("sz", str(sz)), ("space", "4"), ("color", color)):
            e.set(qn(f"w:{k}"), v)
        b.append(e)
        ppr.append(b)

    hp = sec.header.paragraphs[0]
    hp.text = "FitCoach IA — Memoria del proyecto"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs:
        r.font.size, r.font.color.rgb, r.font.name = Pt(9), GREY, "Calibri"
    border(hp, "bottom", "BBBBBB", 4)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Página ")
    r.font.size = Pt(9)
    r = fp.add_run()
    r.font.size = Pt(9)
    field(r, "PAGE")
    r = fp.add_run(" de ")
    r.font.size = Pt(9)
    r = fp.add_run()
    r.font.size = Pt(9)
    field(r, "NUMPAGES")

    # Portada: espacio superior, linea decorativa y salto de pagina tras el autor
    body = d.paragraphs
    title = next(p for p in body if p.style.name == "Title")
    title.paragraph_format.space_before = Pt(170)
    border(title, "bottom", "1A3A6B", 18)
    authors = [p for p in body if p.style.name == "Author"]
    authors[0].paragraph_format.space_before = Pt(60)
    last = authors[-1]
    last.add_run().add_break(WD_BREAK.PAGE)

    # Salto tras el indice
    for p in d.paragraphs:
        if p.style.name == "Heading 1":
            p.paragraph_format.page_break_before = False
            break
    for p in d.paragraphs:
        if (
            p.style.name.startswith("TOC Heading")
            or p.text.strip() == "Table of contents"
            or p.text.strip() == "Índice"
        ):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    first_h1 = next(p for p in d.paragraphs if p.style.name == "Heading 1")
    first_h1.paragraph_format.page_break_before = True

    # Tablas
    def shade(cell, hexcolor):
        tcpr = cell._tc.get_or_add_tcPr()
        s = OxmlElement("w:shd")
        s.set(qn("w:val"), "clear")
        s.set(qn("w:color"), "auto")
        s.set(qn("w:fill"), hexcolor)
        tcpr.append(s)

    for t in d.tables:
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblpr = t._tbl.tblPr
        for old in tblpr.findall(qn("w:tblBorders")):
            tblpr.remove(old)
        bd = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{side}")
            for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "BFC7D5")):
                e.set(qn(f"w:{k}"), v)
            bd.append(e)
        tblpr.append(bd)
        cm = OxmlElement("w:tblCellMar")
        for side, w in (("top", 50), ("left", 90), ("bottom", 50), ("right", 90)):
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:w"), str(w))
            e.set(qn("w:type"), "dxa")
            cm.append(e)
        tblpr.append(cm)
        for i, row in enumerate(t.rows):
            trpr = row._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            trpr.append(cs)
            if i == 0:
                h = OxmlElement("w:tblHeader")
                trpr.append(h)
            for c in row.cells:
                if i == 0:
                    shade(c, "1A3A6B")
                elif i % 2 == 0:
                    shade(c, "F2F5FA")
                for p in c.paragraphs:
                    p.paragraph_format.space_after = Pt(1)
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:
                        r.font.size = Pt(9)
                        if i == 0:
                            r.font.bold = True
                            r.font.color.rgb = RGBColor(255, 255, 255)

    # Imagenes centradas
    for p in d.paragraphs:
        if p._p.xpath(".//w:drawing"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True

    # Actualizar campos (indice, paginas) al abrir
    st = d.settings.element
    u = OxmlElement("w:updateFields")
    u.set(qn("w:val"), "true")
    st.append(u)

    cp = d.core_properties
    cp.title, cp.author = "FitCoach IA — Memoria del proyecto", "Equipo FitCoach IA"
    cp.language = "es-ES"


DOCS = Path(__file__).resolve().parents[2] / "docs"


def main() -> None:
    if "--sin-diagramas" not in sys.argv:
        sys.path.insert(0, str(Path(__file__).parent))
        from diagrams import build_diagrams

        (DOCS / "img").mkdir(exist_ok=True)
        build_diagrams(str(DOCS / "img"))
    out = DOCS / "memoria-proyecto.docx"
    subprocess.run(
        [
            pypandoc.get_pandoc_path(),
            str(DOCS / "memoria-proyecto.md"),
            "-o",
            str(out),
            "--toc",
            "--toc-depth=2",
            f"--resource-path={DOCS}",
        ],
        check=True,
    )
    d = Document(out)
    style_docx(d)
    d.save(out)
    print(f"Generado {out}")


if __name__ == "__main__":
    main()
